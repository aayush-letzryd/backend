"""
LetzRyd Unified Rental Architecture - Master Excel Generator
============================================================
Generates the comprehensive Excel Master Architecture workbook directly
from the live PostgreSQL database.

Tabs created:
  1. How Everything Works (Architecture, 5-tier waterfall, and step-by-step
     guide on how to get rent for Vehicle X with Operator Y)
  2. core_rental_plans (Canonical plan catalogue with integer PKs)
  3. rental_rate_slabs (Dynamic reducing slabs & custom operator trip tiers)
  4. rental_custom_partner_plans (233 active partner rate cards)
  5. rental_model_baselines (21 vehicle model fallback baselines)
  6. rental_fee_rules (7 indemnity fee policies & waiver rules)
  7. rental_exceptions (Audit-grade governance container)
  8. daily_rent_log (Sample output ledger showing explicit rule lineage)
"""

import os
import sys
import psycopg2
from psycopg2.extras import RealDictCursor
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding='utf-8')

DB_HOST = os.getenv('DB_HOST', '35.200.196.113')
DB_PORT = os.getenv('DB_PORT', '5432')
DB_NAME = os.getenv('DB_NAME', 'postgres')
DB_USER = os.getenv('DB_USER', 'postgres')
DB_PASS = os.getenv('DB_PASSWORD', r'8S5]U3@L^Xz)\FH}')

# Styles
NAVY_HEADER_FILL = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
TEAL_HEADER_FILL = PatternFill(start_color="0D5C75", end_color="0D5C75", fill_type="solid")
SECTION_HEADER_FILL = PatternFill(start_color="2C5282", end_color="2C5282", fill_type="solid")
HIGHLIGHT_FILL = PatternFill(start_color="EBF8FF", end_color="EBF8FF", fill_type="solid")
ZEBRA_FILL = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
WHITE_FILL = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

FONT_HEADER = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
FONT_TITLE = Font(name="Calibri", size=16, bold=True, color="1B365D")
FONT_SUBTITLE = Font(name="Calibri", size=11, italic=True, color="4A5568")
FONT_SECTION = Font(name="Calibri", size=13, bold=True, color="FFFFFF")
FONT_BOLD = Font(name="Calibri", size=11, bold=True, color="1A202C")
FONT_REGULAR = Font(name="Calibri", size=11, color="2D3748")
FONT_CODE = Font(name="Consolas", size=10, color="2B6CB0")

THIN_BORDER = Border(
    left=Side(style='thin', color='E2E8F0'),
    right=Side(style='thin', color='E2E8F0'),
    top=Side(style='thin', color='E2E8F0'),
    bottom=Side(style='thin', color='E2E8F0')
)
BOTTOM_DOUBLE_BORDER = Border(
    left=Side(style='thin', color='CBD5E0'),
    right=Side(style='thin', color='CBD5E0'),
    top=Side(style='thin', color='CBD5E0'),
    bottom=Side(style='double', color='1B365D')
)

def create_master_workbook():
    print("=" * 70)
    print("      GENERATING MASTER RENTAL ARCHITECTURE EXCEL")
    print("=" * 70)

    conn = psycopg2.connect(
        host=DB_HOST,
        port=DB_PORT,
        dbname=DB_NAME,
        user=DB_USER,
        password=DB_PASS
    )
    cur = conn.cursor(cursor_factory=RealDictCursor)
    wb = openpyxl.Workbook()

    # -------------------------------------------------------------
    # TAB 1: HOW EVERYTHING WORKS
    # -------------------------------------------------------------
    print("1. Creating 'How Everything Works' guide tab...")
    ws_guide = wb.active
    ws_guide.title = "How Everything Works"
    ws_guide.views.sheetView[0].showGridLines = True

    # Title Banner
    ws_guide.cell(row=1, column=1, value="LETZRYD UNIFIED RENTAL ARCHITECTURE - OPERATIONAL & CALCULATION GUIDE").font = FONT_TITLE
    ws_guide.cell(row=2, column=1, value="Complete technical & business reference: 100% Data-Driven Pricing, Clean Integer IDs, and Operator Slabs").font = FONT_SUBTITLE
    ws_guide.row_dimensions[1].height = 25
    ws_guide.row_dimensions[2].height = 18

    guide_content = [
        ("SECTION 1: CORE ARCHITECTURAL PRINCIPLES", [
            ("1. 100% Data-Driven Pricing", "Zero hardcoded numbers in SQL or application code. Every base rate, dynamic trip reducing tier, fallback rate, and indemnity fee waiver is explicitly stored in physical database tables."),
            ("2. Clean Numeric Primary Keys", "All tables use auto-incrementing integer SERIAL PRIMARY KEY (plan_id: 1, 2, 3...). Human-readable identifiers (e.g. BLR_MASTER_OP) are stored cleanly in plan_code."),
            ("3. Custom Operator Slabs", "Operators with bespoke negotiated trip brackets (e.g. 0-54 trips = 870, 55-64 = 840, 65-74 = 790, 75+ = 770) are stored natively in rental_rate_slabs via partner_id."),
            ("4. Transparent Audit Lineage", "The output ledger (daily_rent_log) stores matched_plan_id, matched_slab_id, and matched_custom_plan_id, linking every single calculated rupee back to its source record."),
            ("5. Zero-Lock Concurrency", "All row-level cascading triggers have been dropped. Daily calculations and weekly Hisaab syncs execute as set-based atomic batch procedures orchestrated by pg_cron.")
        ]),
        ("SECTION 2: THE 5-TIER NATIVE DATABASE WATERFALL PRECEDENCE", [
            ("Priority 1: rental_exceptions", "Audit-grade concessions (approved by Management). Overrides all rules for a specific vehicle and/or partner during an approved date window."),
            ("Priority 2: rental_custom_partner_plans", "Negotiated Partner Agreements. 233 active partner rate cards. When a partner has an agreed flat daily rent (e.g. 940/day), it wins here."),
            ("Priority 3: rental_rate_slabs", "Dynamic Reducing Trip Slabs. Checked in 2 sub-tiers: (A) Specific Operator Slabs (partner_id = X); (B) Standard City Slabs (partner_id = 'ALL'). Matches on weekly completed trips."),
            ("Priority 4: rental_model_baselines", "Vehicle Model Rate Fallbacks. If no slabs or custom agreements apply, matches the car model in that city (e.g. Mumbai Dzire: 1,100, Hyderabad eC3: 1,400, WagonR: 989)."),
            ("Priority 5: core_rental_plans", "Master City Fallback Default. The ultimate safety net plan defined in the catalogue for that city (default_daily_rent: BLR 929, HYD 989, MUM 970)."),
            ("Fee Rules: rental_fee_rules", "Indemnity fees are evaluated in parallel: Partner Specific -> Vehicle Model Specific (e.g. Xcent = 0) -> City Specific (e.g. Mumbai = 0) -> Global Default (30/day).")
        ]),
        ("SECTION 3: HOW TO USE THE TABLES TO GET RENT FOR VEHICLE 'X' WITH OPERATOR 'Y'", [
            ("STEP 1: Identify Key Attributes", "Collect: Vehicle Number (X), Operator / Partner Code (Y), Operational City, Car Model, Attendance Status (e.g. Active, Maintenance), and Weekly Completed Trips (Uber + Ola)."),
            ("STEP 2: Check Billability", "If Attendance Status in ('Drop Off', 'Drop-off', 'RFD', 'Unassigned') and Weekly Trips == 0 -> Non-billable: Rent = ₹0, Indemnity = ₹0. If Maintenance without billable flag and 0 trips -> Rent = ₹0."),
            ("STEP 3: Check Priority 1 (Exceptions)", "Query rental_exceptions WHERE partner_id = Y OR vehicle_number = X AND valid_from <= date <= valid_to. If found, apply override rent & fee."),
            ("STEP 4: Check Priority 2 (Custom Flat Agreement)", "Query rental_custom_partner_plans WHERE partner_id = Y AND is_active = TRUE. If custom_daily_rent IS NOT NULL, apply flat rent (e.g. ₹940/day)."),
            ("STEP 5: Check Priority 3 (Dynamic Trip Slabs)", "Query rental_rate_slabs WHERE city = City AND (partner_id = Y OR partner_id = 'ALL') AND (customer_type = 'Operator' OR customer_type = 'ALL') AND trip_min <= Trips <= trip_max. If partner_id = Y exists, IT WINS over 'ALL'!"),
            ("STEP 6: Check Priority 4 (Model Baseline)", "If no slab matched, query rental_model_baselines WHERE city = City AND vehicle_model = Model. Apply default_base_rent."),
            ("STEP 7: Check Priority 5 (City Fallback)", "If model not found, query core_rental_plans WHERE city = City AND calculation_type = 'MODEL_FALLBACK'. Apply default_daily_rent."),
            ("STEP 8: Look up Indemnity Fee", "Query rental_fee_rules for the best match: Specific Partner -> Specific Model -> Specific City -> Global Default. If is_waiver = TRUE, Fee = ₹0."),
            ("STEP 9: Calculate Net Rent", "Net Daily Rent = Applied Daily Rent + Applied Daily Indemnity. Record in daily_rent_log with matched IDs.")
        ]),
        ("SECTION 4: WORKED NUMERICAL EXAMPLES (OPERATOR & DRIVER SCENARIOS)", [
            ("Scenario A: Operator with Custom Tiered Deal", "Vehicle KA05AQ1234 under Operator 'LETZBLR_HAMZA'. Car does 62 trips this week. Engine checks rental_rate_slabs: Matches Slab #49 (Partner: LETZBLR_HAMZA, Min: 55, Max: 64) -> Rent = ₹840.00/day. Fee = ₹30.00. Net = ₹870.00/day."),
            ("Scenario B: Operator with High Trips", "Same Operator 'LETZBLR_HAMZA', but car achieves 82 trips this week. Engine matches Slab #51 (Partner: LETZBLR_HAMZA, Min: 75, Max: 9999) -> Rent = ₹770.00/day. Fee = ₹30.00. Net = ₹800.00/day."),
            ("Scenario C: Standard Operator (No Custom Deal)", "Vehicle KA05AQ9999 under Operator 'LETZBLRIP7025077468' (no custom card). Car does 115 trips in Bangalore. Engine matches standard Operator Plan #2 (BLR_MASTER_OP), Slab #42 (Min: 110, Max: 129) -> Rent = ₹500.00/day. Fee = ₹30.00. Net = ₹530.00/day."),
            ("Scenario D: Mumbai Vehicle (City Fee Waiver)", "Vehicle MH03FC5592 under Operator 'LETZMUMIP9167567426'. Car does 104 trips in Mumbai. Engine matches Mumbai Reducing Plan #10, Slab #30 (Min: 95, Max: 109) -> Rent = ₹659.00/day. rental_fee_rules matches Mumbai Waiver -> Fee = ₹0.00. Net = ₹659.00/day."),
            ("Scenario E: Retired Fleet Vehicle (Hyderabad Xcent)", "Vehicle TS07UB1111 (Hyundai Xcent) under any partner. rental_model_baselines gives base rent ₹900.00/day. rental_fee_rules has Xcent waiver -> Fee = ₹0.00. Net = ₹900.00/day.")
        ]),
        ("SECTION 5: COPY-PASTE SQL QUERY TO GET RENT FOR VEHICLE 'X' WITH OPERATOR 'Y'", [
            ("SQL Stored Procedure Call", "CALL public.sp_calculate_daily_rent('2026-09-20'::date, '2026-09-20'::date); -- Calculates all vehicles and outputs into daily_rent_log with lineage"),
            ("Direct SQL Verification Query", """SELECT log_date, vehicle_number, partner_id, city, vehicle_model, weekly_completed_trips, applied_daily_rent, applied_daily_indemnity, net_daily_rent, matched_plan_id, matched_slab_id, matched_custom_plan_id, calculation_rule FROM public.daily_rent_log WHERE vehicle_number = 'KA05AQ1234' AND log_date = '2026-09-20';""")
        ])
    ]

    curr_row = 4
    for section_title, items in guide_content:
        ws_guide.cell(row=curr_row, column=1, value=section_title).font = FONT_SECTION
        ws_guide.cell(row=curr_row, column=1).fill = SECTION_HEADER_FILL
        ws_guide.merge_cells(start_row=curr_row, start_column=1, end_row=curr_row, end_column=3)
        ws_guide.row_dimensions[curr_row].height = 24
        curr_row += 1

        for label, desc in items:
            ws_guide.cell(row=curr_row, column=1, value=label).font = FONT_BOLD
            ws_guide.cell(row=curr_row, column=1).alignment = Alignment(vertical='top')
            ws_guide.cell(row=curr_row, column=2, value=desc).font = FONT_REGULAR
            ws_guide.cell(row=curr_row, column=2).alignment = Alignment(wrap_text=True, vertical='top')
            ws_guide.merge_cells(start_row=curr_row, start_column=2, end_row=curr_row, end_column=3)
            ws_guide.row_dimensions[curr_row].height = 28 if len(desc) > 80 else 20
            curr_row += 1
        curr_row += 1

    ws_guide.column_dimensions['A'].width = 38
    ws_guide.column_dimensions['B'].width = 55
    ws_guide.column_dimensions['C'].width = 45

    # -------------------------------------------------------------
    # HELPER: EXPORT DATABASE TABLE TO SHEET
    # -------------------------------------------------------------
    def export_table_to_sheet(table_name, query, sheet_title, col_headers=None, num_formats=None):
        print(f"Exporting '{sheet_title}'...")
        ws = wb.create_sheet(title=sheet_title)
        ws.views.sheetView[0].showGridLines = True

        cur.execute(query)
        rows = cur.fetchall()
        cols = [desc[0] for desc in cur.description] if not col_headers else list(col_headers.keys())

        # Header Row
        ws.row_dimensions[1].height = 26
        for col_idx, col_name in enumerate(cols, 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.value = col_headers.get(col_name, col_name.replace('_', ' ').title()) if col_headers else col_name.replace('_', ' ').title()
            cell.font = FONT_HEADER
            cell.fill = NAVY_HEADER_FILL
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = THIN_BORDER

        # Data Rows
        for row_idx, r in enumerate(rows, 2):
            ws.row_dimensions[row_idx].height = 20
            fill = ZEBRA_FILL if row_idx % 2 == 0 else WHITE_FILL
            for col_idx, col_name in enumerate(cols, 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                val = r[col_name]
                cell.value = val
                cell.font = FONT_REGULAR
                cell.fill = fill
                cell.border = THIN_BORDER

                # Formatting
                fmt = num_formats.get(col_name) if num_formats else None
                if fmt == 'currency':
                    cell.number_format = '₹#,##0.00'
                    cell.alignment = Alignment(horizontal='right', vertical='center')
                elif fmt == 'integer':
                    cell.number_format = '#,##0'
                    cell.alignment = Alignment(horizontal='right', vertical='center')
                elif fmt == 'date':
                    cell.number_format = 'YYYY-MM-DD'
                    cell.alignment = Alignment(horizontal='center', vertical='center')
                elif fmt == 'center':
                    cell.alignment = Alignment(horizontal='center', vertical='center')
                else:
                    cell.alignment = Alignment(horizontal='left', vertical='center')

        # Auto-adjust column widths
        for col in ws.columns:
            col_letter = get_column_letter(col[0].column)
            max_len = max(len(str(cell.value or '')) for cell in col[:100])
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)
        ws.freeze_panes = "A2"

    # -------------------------------------------------------------
    # TAB 2: core_rental_plans
    # -------------------------------------------------------------
    export_table_to_sheet(
        "core_rental_plans",
        """
        SELECT 
            plan_id, plan_code, city, plan_name, plan_category, calculation_type,
            default_daily_rent, default_daily_fee, description, is_active
        FROM public.core_rental_plans
        ORDER BY plan_id;
        """,
        "core_rental_plans",
        num_formats={
            'plan_id': 'integer',
            'default_daily_rent': 'currency',
            'default_daily_fee': 'currency',
            'is_active': 'center'
        }
    )

    # -------------------------------------------------------------
    # TAB 3: rental_rate_slabs
    # -------------------------------------------------------------
    export_table_to_sheet(
        "rental_rate_slabs",
        """
        SELECT 
            s.slab_id, s.plan_id, p.plan_code, s.partner_id, s.city, s.customer_type,
            s.vehicle_model, s.metric_type, s.condition_rule, s.trip_min, s.trip_max,
            s.base_daily_rent, s.default_daily_fee, s.evidence_reference
        FROM public.rental_rate_slabs s
        JOIN public.core_rental_plans p ON p.plan_id = s.plan_id
        ORDER BY s.plan_id, s.partner_id, s.trip_min;
        """,
        "rental_rate_slabs",
        num_formats={
            'slab_id': 'integer',
            'plan_id': 'integer',
            'trip_min': 'integer',
            'trip_max': 'integer',
            'base_daily_rent': 'currency',
            'default_daily_fee': 'currency'
        }
    )

    # -------------------------------------------------------------
    # TAB 4: rental_custom_partner_plans
    # -------------------------------------------------------------
    export_table_to_sheet(
        "rental_custom_partner_plans",
        """
        SELECT 
            custom_plan_id, partner_id, partner_name, city, vehicle_model,
            vehicle_number, plan_id, custom_daily_rent, custom_daily_fee,
            plan_label, evidence_source, approved_by, valid_from, valid_to, is_active
        FROM public.rental_custom_partner_plans
        ORDER BY custom_plan_id;
        """,
        "rental_custom_partner_plans",
        num_formats={
            'custom_plan_id': 'integer',
            'plan_id': 'integer',
            'custom_daily_rent': 'currency',
            'custom_daily_fee': 'currency',
            'valid_from': 'date',
            'valid_to': 'date',
            'is_active': 'center'
        }
    )

    # -------------------------------------------------------------
    # TAB 5: rental_model_baselines
    # -------------------------------------------------------------
    export_table_to_sheet(
        "rental_model_baselines",
        """
        SELECT 
            baseline_id, city, vehicle_model, default_base_rent,
            default_daily_indemnity, all_platform_flat_rent, is_active
        FROM public.rental_model_baselines
        ORDER BY city, baseline_id;
        """,
        "rental_model_baselines",
        num_formats={
            'baseline_id': 'integer',
            'default_base_rent': 'currency',
            'default_daily_indemnity': 'currency',
            'all_platform_flat_rent': 'currency',
            'is_active': 'center'
        }
    )

    # -------------------------------------------------------------
    # TAB 6: rental_fee_rules
    # -------------------------------------------------------------
    export_table_to_sheet(
        "rental_fee_rules",
        """
        SELECT 
            fee_rule_id, city, partner_id, vehicle_model,
            fee_amount, is_waiver, reason, valid_from, valid_to
        FROM public.rental_fee_rules
        ORDER BY fee_rule_id;
        """,
        "rental_fee_rules",
        num_formats={
            'fee_rule_id': 'integer',
            'fee_amount': 'currency',
            'is_waiver': 'center',
            'valid_from': 'date',
            'valid_to': 'date'
        }
    )

    # -------------------------------------------------------------
    # TAB 7: rental_exceptions
    # -------------------------------------------------------------
    export_table_to_sheet(
        "rental_exceptions",
        """
        SELECT 
            exception_id, override_type, city, partner_id, vehicle_number,
            vehicle_model, override_daily_rent, override_fee, canonical_expected_rent,
            variance, reason, status, approved_by, valid_from, valid_to
        FROM public.rental_exceptions
        ORDER BY exception_id;
        """,
        "rental_exceptions",
        num_formats={
            'exception_id': 'integer',
            'override_daily_rent': 'currency',
            'override_fee': 'currency',
            'canonical_expected_rent': 'currency',
            'variance': 'currency',
            'valid_from': 'date',
            'valid_to': 'date'
        }
    )

    # -------------------------------------------------------------
    # TAB 8: daily_rent_log (Sample Live Calculations)
    # -------------------------------------------------------------
    export_table_to_sheet(
        "daily_rent_log",
        """
        SELECT 
            id, log_date, week_id, vehicle_number, partner_id, city,
            vehicle_model, attendance_status, is_billable_day, weekly_completed_trips,
            applied_daily_rent, applied_daily_indemnity, net_daily_rent,
            matched_plan_id, matched_slab_id, matched_custom_plan_id,
            calculation_rule
        FROM public.daily_rent_log
        WHERE log_date >= (CURRENT_DATE - INTERVAL '14 days')
        ORDER BY log_date DESC, id DESC
        LIMIT 500;
        """,
        "daily_rent_log",
        num_formats={
            'id': 'integer',
            'log_date': 'date',
            'is_billable_day': 'center',
            'weekly_completed_trips': 'integer',
            'applied_daily_rent': 'currency',
            'applied_daily_indemnity': 'currency',
            'net_daily_rent': 'currency',
            'matched_plan_id': 'integer',
            'matched_slab_id': 'integer',
            'matched_custom_plan_id': 'integer'
        }
    )

    # Save to file
    out_dir = r"c:\Users\anura\Downloads\LetzRyd Antigravity\backend_repo\Rental Final Table"
    path1 = os.path.join(out_dir, "LetzRyd_Rental_Master_Architecture.xlsx")
    path2 = os.path.join(out_dir, "LetzRyd_Rental_Master_Architecture_Complete.xlsx")

    for p in [path1, path2]:
        try:
            print(f"Saving to: {p}")
            wb.save(p)
            print(f"  [OK] Saved successfully: {os.path.basename(p)}")
        except PermissionError:
            alt_path = os.path.join(out_dir, os.path.splitext(os.path.basename(p))[0] + "_v2.xlsx")
            print(f"  [LOCKED] File is open in Excel: {os.path.basename(p)}. Saving to: {os.path.basename(alt_path)}")
            wb.save(alt_path)
            print(f"  [OK] Saved successfully: {os.path.basename(alt_path)}")

    cur.close()
    conn.close()
    print("\n" + "=" * 70)
    print("MASTER WORKBOOK GENERATED SUCCESSFULLY!")
    print(f"Total Sheets: {len(wb.sheetnames)}")
    print(f"Sheets: {wb.sheetnames}")
    print("=" * 70)

if __name__ == '__main__':
    create_master_workbook()
