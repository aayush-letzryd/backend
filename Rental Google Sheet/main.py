import os
import io
import requests
import openpyxl
import psycopg2
import functions_framework

def get_db_connection():
    return psycopg2.connect(
        host=os.getenv("DB_HOST", "35.200.196.113"),
        port=os.getenv("DB_PORT", "5432"),
        dbname=os.getenv("DB_NAME", "postgres"),
        user=os.getenv("DB_USER", "postgres"),
        password=os.getenv("DB_PASSWORD", r"8S5]U3@L^Xz)\FH}")
    )

def ensure_tables(cur):
    ddl = """
    CREATE TABLE IF NOT EXISTS sheet_rental_slabs (
        id SERIAL PRIMARY KEY,
        city VARCHAR(32) NOT NULL,
        vehicle_model VARCHAR(64) NOT NULL,
        uber_type VARCHAR(32) DEFAULT 'TBS',
        plan_scheme VARCHAR(64) NOT NULL DEFAULT 'Uber Reducing Rent',
        trip_slab_label VARCHAR(64) NOT NULL,
        min_trips INT NOT NULL DEFAULT 0,
        max_trips INT NOT NULL DEFAULT 9999,
        daily_rent NUMERIC(10,2) NOT NULL,
        daily_indemnity NUMERIC(10,2) NOT NULL DEFAULT 30.00,
        platform VARCHAR(64) DEFAULT 'Uber',
        pass_on_incentive VARCHAR(16) DEFAULT 'yes',
        last_synced_at TIMESTAMP WITHOUT TIME ZONE DEFAULT CURRENT_TIMESTAMP,
        CONSTRAINT uq_sheet_rental_slabs UNIQUE (city, vehicle_model, uber_type, plan_scheme, min_trips)
    );

    CREATE TABLE IF NOT EXISTS sheet_rental_partners (
        vendor_code VARCHAR(64) PRIMARY KEY,
        vendor_name VARCHAR(128),
        city VARCHAR(32) NOT NULL,
        vendor_type VARCHAR(32),
        plan_name VARCHAR(64),
        platform VARCHAR(64),
        plan_type_hisaab VARCHAR(64),
        custom_daily_rent NUMERIC(10,2) NULL,
        custom_daily_indemnity NUMERIC(10,2) NULL,
        last_synced_at TIMESTAMP WITHOUT TIME ZONE DEFAULT CURRENT_TIMESTAMP
    );
    """
    cur.execute(ddl)

def sync_rental_data():
    sheet_id = os.getenv("SPREADSHEET_ID", "1xnGg3qhb1AnCP2Qv6e2gmd0zCi5j-bNbx9yDc7etzf4")
    export_url = f"https://docs.google.com/spreadsheets/d/{sheet_id}/export?format=xlsx"
    
    resp = requests.get(export_url, headers={"User-Agent": "Mozilla/5.0"}, timeout=45)
    resp.raise_for_status()
    wb = openpyxl.load_workbook(io.BytesIO(resp.content), data_only=True)
    
    conn = get_db_connection()
    cur = conn.cursor()
    ensure_tables(cur)
    
    # --- 1. sheet_rental_partners ---
    partners = []
    
    # HYD
    if "HYD - Driver platform" in wb.sheetnames:
        ws = wb["HYD - Driver platform"]
        for r in list(ws.iter_rows(values_only=True))[1:]:
            if not r or not r[1]: continue
            vname = str(r[0]).strip() if r[0] else ""
            vcode = str(r[1]).strip()
            vtype = str(r[2]).strip() if len(r) > 2 and r[2] else "Operator"
            plan = str(r[3]).strip() if len(r) > 3 and r[3] else ""
            plat = str(r[4]).strip() if len(r) > 4 and r[4] else ""
            ptype = str(r[5]).strip() if len(r) > 5 and r[5] else ""
            
            custom_rent, custom_indem = None, None
            if "900" in vname or "900" in plan: custom_rent = 900.00
            if "970" in vname or "970" in plan: custom_rent = 970.00
            if vcode == "LETZHYDIP9885838038": custom_rent, custom_indem = 900.00, 0.00
            elif vcode == "LETZHYDIP9848529242": custom_rent = 900.00
            elif vcode == "LETZHYDIP7396655106": custom_rent = 970.00
            elif "All Platform" in ptype or "All Platform" in plat: custom_rent = 1050.00

            partners.append((vcode, vname, "Hyderabad", vtype, plan, plat, ptype, custom_rent, custom_indem))

    # MUM
    if "MUM - Driver platform" in wb.sheetnames:
        ws = wb["MUM - Driver platform"]
        for r in list(ws.iter_rows(values_only=True))[1:]:
            if not r or not r[1]: continue
            vname = str(r[0]).strip() if r[0] else ""
            vcode = str(r[1]).strip()
            plan = str(r[2]).strip() if len(r) > 2 and r[2] else ""
            ptype = str(r[3]).strip() if len(r) > 3 and r[3] else ""
            custom_rent = 1050.00 if ("All Platform" in ptype or "All Platform" in plan) else None
            partners.append((vcode, vname, "Mumbai", "Operator", plan, "Uber", ptype, custom_rent, None))

    # BLR
    if "BLR - Driver platform" in wb.sheetnames:
        ws = wb["BLR - Driver platform"]
        for r in list(ws.iter_rows(values_only=True))[1:]:
            if not r or not r[1]: continue
            vname = str(r[0]).strip() if r[0] else ""
            vcode = str(r[1]).strip()
            vtype = str(r[2]).strip() if len(r) > 2 and r[2] else "Individual"
            ptype = str(r[3]).strip() if len(r) > 3 and r[3] else ""
            custom_rent = 1050.00 if ("All Platform" in ptype) else None
            custom_indem = 15.00 if vcode == "LETZBLRIP9035252877" else None
            partners.append((vcode, vname, "Bengaluru", vtype, "", "Uber", ptype, custom_rent, custom_indem))

    upsert_partner_sql = """
    INSERT INTO sheet_rental_partners (
        vendor_code, vendor_name, city, vendor_type, plan_name, platform, plan_type_hisaab, custom_daily_rent, custom_daily_indemnity, last_synced_at
    ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, CURRENT_TIMESTAMP)
    ON CONFLICT (vendor_code) DO UPDATE SET
        vendor_name = EXCLUDED.vendor_name,
        city = EXCLUDED.city,
        vendor_type = EXCLUDED.vendor_type,
        plan_name = EXCLUDED.plan_name,
        platform = EXCLUDED.platform,
        plan_type_hisaab = EXCLUDED.plan_type_hisaab,
        custom_daily_rent = EXCLUDED.custom_daily_rent,
        custom_daily_indemnity = EXCLUDED.custom_daily_indemnity,
        last_synced_at = CURRENT_TIMESTAMP;
    """
    for p in partners:
        cur.execute(upsert_partner_sql, p)

    # --- 2. sheet_rental_slabs ---
    slabs = []

    # HYD Slabs
    if "HYD - Rental Slab" in wb.sheetnames:
        ws = wb["HYD - Rental Slab"]
        for r in list(ws.iter_rows(values_only=True))[1:]:
            if not r or not r[0] or not r[3]: continue
            veh = str(r[0]).strip()
            try: min_t = int(float(r[1])) if r[1] is not None else 0
            except: min_t = 0
            try: rent = float(r[3])
            except: continue
            utype = str(r[4]).strip() if len(r) > 4 and r[4] else "TBS"
            max_t = 9999 if min_t >= 140 else (min_t + 19 if min_t > 0 else 49)
            slab_label = f"{min_t}-{max_t}" if max_t < 9999 else f"{min_t}+"
            slabs.append(("Hyderabad", veh, utype, "Uber Reducing Rent", slab_label, min_t, max_t, rent, 30.00, "Uber", "yes"))

    # MUM Slabs
    if "MUM - Rental Slab" in wb.sheetnames:
        ws = wb["MUM - Rental Slab"]
        for r in list(ws.iter_rows(values_only=True))[1:]:
            if not r or r[0] is None or r[1] is None: continue
            try: min_t = int(float(r[0])); rent = float(r[1])
            except: continue
            if min_t == 0: max_t, slab_label = 64, "0-64"
            elif min_t == 65: max_t, slab_label = 79, "65-79"
            elif min_t == 80: max_t, slab_label = 99, "80-99"
            elif min_t == 100: max_t, slab_label = 119, "100-119"
            else: max_t, slab_label = 9999, f"{min_t}+"
            slabs.append(("Mumbai", "Maruti Wagonr Tour H3 CNG", "TBS", "Uber Reducing Rent", slab_label, min_t, max_t, rent, 30.00, "Uber", "yes"))

    # BLR Slabs
    if "BLR- Rental Slab" in wb.sheetnames:
        ws = wb["BLR- Rental Slab"]
        for r in list(ws.iter_rows(values_only=True))[1:]:
            if not r or not r[0] or not r[2]: continue
            utype = str(r[0]).strip()
            slab_label = str(r[1]).strip()
            try: rent = float(r[2])
            except: continue
            if "<55" in slab_label: min_t, max_t = 0, 54
            elif "55+" in slab_label: min_t, max_t = 55, 9999
            elif "65+" in slab_label: min_t, max_t = 65, 9999
            elif "70+" in slab_label: min_t, max_t = 70, 9999
            elif "75+" in slab_label: min_t, max_t = 75, 9999
            elif "90+" in slab_label: min_t, max_t = 90, 9999
            else: min_t, max_t = 0, 9999
            slabs.append(("Bengaluru", "Maruti Wagonr Tour H3 CNG", utype, "Uber Reducing Rent", slab_label, min_t, max_t, rent, 30.00, "Uber", "yes"))

    upsert_slab_sql = """
    INSERT INTO sheet_rental_slabs (
        city, vehicle_model, uber_type, plan_scheme, trip_slab_label, min_trips, max_trips, daily_rent, daily_indemnity, platform, pass_on_incentive, last_synced_at
    ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, CURRENT_TIMESTAMP)
    ON CONFLICT (city, vehicle_model, uber_type, plan_scheme, min_trips) DO UPDATE SET
        trip_slab_label = EXCLUDED.trip_slab_label,
        max_trips = EXCLUDED.max_trips,
        daily_rent = EXCLUDED.daily_rent,
        daily_indemnity = EXCLUDED.daily_indemnity,
        platform = EXCLUDED.platform,
        pass_on_incentive = EXCLUDED.pass_on_incentive,
        last_synced_at = CURRENT_TIMESTAMP;
    """
    for s in slabs:
        cur.execute(upsert_slab_sql, s)

    conn.commit()
    cur.close()
    conn.close()
    
    return {
        "status": "success",
        "partners_synced": len(partners),
        "slabs_synced": len(slabs)
    }

@functions_framework.http
def sync_rental_sheets_http(request):
    try:
        res = sync_rental_data()
        return res, 200
    except Exception as e:
        return {"status": "error", "message": str(e)}, 500

if __name__ == "__main__":
    print("Running local synchronization test...")
    res = sync_rental_data()
    print("Sync result:", res)
